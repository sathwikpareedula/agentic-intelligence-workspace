"""Deterministic transform-to-template planning, execution, and validation."""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
from typing import Any
from uuid import uuid4

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

from app.models.template_transforms import (
    ClarificationRequirement,
    DerivationRule,
    FieldMapping,
    FieldProvenance,
    SourceInspection,
    SourcePayload,
    TemplateColumn,
    TemplateInspection,
    TemplateSheet,
    TransformExecutionRequest,
    TransformProposal,
    TransformProposalRequest,
    TransformTemplatePlan,
    ValidationCheck,
    OutputValidation,
)
from app.models.transformations import JoinSpec
from app.services.artifacts import GeneratedArtifact
from app.services.datasets import MAX_UPLOAD_BYTES, inspect_dataset, load_dataset
from app.services.transformations import TransformationError, join_datasets
from app.services.workbook_safety import (
    WorkbookArchiveError,
    WorkbookArchiveTooLargeError,
    validate_workbook_archive,
)


MAX_TEMPLATE_COLUMNS = 500
MAX_TEMPLATE_ROWS = 100_000
MAX_TEMPLATE_CELLS = 250_000
FORMULA_PREFIXES = ("=", "+", "-", "@")
DOCUMENTED_ALIASES = {
    "customerid": {"customer", "customerkey", "custid", "custkey", "accountid"},
    "firstname": {"givenname", "forename"},
    "lastname": {"surname", "familyname"},
    "fullname": {"customername", "employeename", "salespersonname"},
    "region": {"territory", "salesarea", "marketregion"},
    "grosssales": {"grossamount", "salesamount", "amount"},
    "netsales": {"netamount", "eligibleamount"},
    "returns": {"returnamount", "refunds", "refundamount"},
    "orderdate": {"transactiondate", "purchasedate", "date"},
}


class TemplateTransformError(Exception):
    """Base class for safe, user-facing transform failures."""


class TemplateInspectionError(TemplateTransformError):
    """Raised when a target template cannot be inspected safely."""


class ClarificationRequiredError(TemplateTransformError):
    def __init__(self, clarifications: list[ClarificationRequirement]) -> None:
        super().__init__("Required target fields remain unresolved.")
        self.clarifications = clarifications


class TransformValidationError(TemplateTransformError):
    """Raised when execution or output validation fails closed."""


@dataclass(frozen=True)
class ExecutedTemplateTransform:
    artifact: GeneratedArtifact
    plan: TransformTemplatePlan
    validation: OutputValidation
    provenance: list[FieldProvenance]


def inspect_template(filename: str, content: bytes, sheet: str | None = None, header_row: int | None = None) -> TemplateInspection:
    extension = Path(filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise TemplateInspectionError("Target template must be a .csv or .xlsx file; macro-enabled workbooks are not accepted.")
    if not content:
        raise TemplateInspectionError("The target template is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise TemplateInspectionError(f"Target template exceeds the {MAX_UPLOAD_BYTES // (1024 * 1024)} MiB limit.")
    if extension == ".csv":
        return _inspect_csv_template(filename, content)
    _validate_workbook_archive(content)
    try:
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=False, keep_links=False)
    except Exception as exc:
        raise TemplateInspectionError("Could not safely read the XLSX target template.") from exc
    if not workbook.sheetnames:
        raise TemplateInspectionError("The target workbook contains no worksheets.")
    selected = sheet or next((name for name in workbook.sheetnames if workbook[name].sheet_state == "visible"), workbook.sheetnames[0])
    if selected not in workbook.sheetnames:
        raise TemplateInspectionError(f"Target sheet '{selected}' was not found.")
    worksheet = workbook[selected]
    if (
        worksheet.max_column > MAX_TEMPLATE_COLUMNS
        or worksheet.max_row > MAX_TEMPLATE_ROWS
        or worksheet.max_column * worksheet.max_row > MAX_TEMPLATE_CELLS
    ):
        raise TemplateInspectionError("The target worksheet exceeds supported dimensions.")
    selected_header_row = header_row or _detect_header_row(worksheet)
    headers = _headers_from_row(worksheet, selected_header_row)
    columns = [_inspect_template_column(worksheet, selected_header_row, index, name) for index, name in enumerate(headers, 1)]
    formula_count = 0
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                _validate_template_formula(str(cell.value))
                formula_count += 1
    sheets = [
        TemplateSheet(name=item.title, state=item.sheet_state, max_row=max(1, item.max_row), max_column=max(1, item.max_column))
        for item in workbook.worksheets
    ]
    fingerprint = _template_fingerprint("xlsx", selected, selected_header_row, headers, columns, sheets)
    return TemplateInspection(
        filename=Path(filename).name,
        file_type="xlsx",
        target_sheet=selected,
        header_row=selected_header_row,
        headers=headers,
        columns=columns,
        sheets=sheets,
        formula_cell_count=formula_count,
        has_macros=False,
        fingerprint=fingerprint,
    )


def propose_transform(request: TransformProposalRequest) -> TransformProposal:
    _validate_unique_roles(request.sources)
    template = inspect_template(
        request.target.filename,
        request.target.content(),
        request.target.sheet,
        request.header_row,
    )
    source_inspections = [_inspect_source(source) for source in request.sources]
    source_by_role = {item.role: item for item in request.sources}
    inspection_by_role = {item.role: item for item in source_inspections}
    formula_targets = {column.name for column in template.columns if column.formula is not None}
    required_fields = request.required_fields or [name for name in template.headers if name not in formula_targets]
    unknown_required = sorted(set(required_fields) - set(template.headers))
    if unknown_required:
        raise TemplateInspectionError(f"Required fields are not present in the target template: {unknown_required}.")
    mappings = [
        _propose_mapping(
            column,
            source_by_role,
            inspection_by_role,
            request.explicit_mappings.get(column.name),
            request.documented_aliases.get(column.name, []),
        )
        for column in template.columns
    ]
    clarifications = _clarifications(mappings, required_fields, set())
    plan = TransformTemplatePlan(
        source_roles=[source.role for source in request.sources],
        expected_source_columns={item.role: item.inspection.columns for item in source_inspections},
        target_filename=template.filename,
        target_sheet=template.target_sheet,
        target_header_row=template.header_row,
        target_headers=template.headers,
        template_fingerprint=template.fingerprint,
        mappings=mappings,
        required_fields=required_fields,
        unique_fields=[],
    )
    return TransformProposal(
        status="clarification_required" if clarifications else "ready",
        template=template,
        sources=source_inspections,
        plan=plan,
        clarifications=clarifications,
    )


def execute_transform(request: TransformExecutionRequest) -> ExecutedTemplateTransform:
    _validate_unique_roles(request.sources)
    plan = request.plan
    roles = [source.role for source in request.sources]
    if roles != plan.source_roles:
        raise TransformValidationError(f"Source-role drift detected: expected {plan.source_roles}, received {roles}.")
    current_template = inspect_template(
        request.target.filename,
        request.target.content(),
        plan.target_sheet,
        plan.target_header_row,
    )
    if current_template.fingerprint != plan.template_fingerprint:
        raise TransformValidationError("Template drift detected; inspect and confirm mappings against the current template before rerunning.")
    if current_template.headers != plan.target_headers:
        raise TransformValidationError("Target schema/order drift detected.")

    loaded = {}
    for source in request.sources:
        dataset = load_dataset(source.filename, source.content(), source.sheet)
        actual = inspect_dataset(dataset).columns
        expected = plan.expected_source_columns.get(source.role)
        if expected != actual:
            raise TransformValidationError(
                f"Schema drift detected for {source.role}: expected columns {expected}, received {actual}."
            )
        loaded[source.role] = dataset.frame.copy()

    derivation_targets = {rule.target_field for rule in plan.derivations}
    clarifications = _clarifications(plan.mappings, plan.required_fields, derivation_targets)
    if clarifications:
        raise ClarificationRequiredError(clarifications)
    _validate_plan_targets(plan)

    merged, join_diagnostics = _join_sources(plan, loaded)
    output = pd.DataFrame(index=merged.index)
    provenance: list[FieldProvenance] = []
    mapping_by_target = {mapping.target_field: mapping for mapping in plan.mappings}
    column_by_target = {column.name: column for column in current_template.columns}

    for target in plan.target_headers:
        mapping = mapping_by_target[target]
        if mapping.mapping_type == "template_formula" or target in derivation_targets:
            output[target] = None
            continue
        if mapping.source_role is None or mapping.source_field is None:
            output[target] = None
            continue
        series = _source_series(merged, mapping.source_role, mapping.source_field)
        output[target] = _apply_mapping_transformation(series, mapping.transformation)
        output[target] = _coerce_target_type(output[target], column_by_target[target].inferred_type, target)
        provenance.append(
            FieldProvenance(
                target_field=target,
                source_fields=[f"{mapping.source_role}.{mapping.source_field}"],
                transformation=mapping.transformation,
                validation="passed",
            )
        )

    for rule in plan.derivations:
        series, evidence_ids = _derive(rule, merged, output, request.policy_evidence)
        output[rule.target_field] = _coerce_target_type(series, column_by_target[rule.target_field].inferred_type, rule.target_field)
        provenance.append(
            FieldProvenance(
                target_field=rule.target_field,
                source_fields=[f"{item.source_role}.{item.source_field}" for item in rule.inputs],
                transformation=rule.operation,
                policy_evidence_ids=evidence_ids,
                validation="passed",
            )
        )

    for target in plan.target_headers:
        mapping = mapping_by_target[target]
        if mapping.mapping_type == "template_formula":
            provenance.append(
                FieldProvenance(
                    target_field=target,
                    source_fields=[],
                    transformation="trusted_template_formula",
                    validation="passed",
                )
            )

    errors, warnings, checks = _validate_output_frame(output, plan, join_diagnostics)
    if errors:
        raise TransformValidationError("; ".join(errors))
    artifact = _write_artifact(request, current_template, output, provenance)
    reopen_checks = _validate_artifact_reopens(artifact, current_template, len(output))
    checks.extend(reopen_checks)
    reopen_errors = [check.detail for check in reopen_checks if not check.passed]
    if reopen_errors:
        raise TransformValidationError("; ".join(reopen_errors))
    status = "completed_with_warnings" if warnings else "completed"
    validation = OutputValidation(
        status=status,
        checks=checks,
        errors=[],
        warnings=warnings,
        input_row_count=len(loaded[plan.source_roles[0]]),
        output_row_count=len(output),
        join_diagnostics=join_diagnostics,
    )
    return ExecutedTemplateTransform(artifact=artifact, plan=plan, validation=validation, provenance=provenance)


def _inspect_csv_template(filename: str, content: bytes) -> TemplateInspection:
    try:
        frame = pd.read_csv(BytesIO(content), encoding="utf-8-sig", nrows=5, on_bad_lines="error")
    except Exception as exc:
        raise TemplateInspectionError("Could not read the CSV target template.") from exc
    headers = [str(column) for column in frame.columns]
    _validate_headers(headers)
    columns = []
    for index, name in enumerate(headers, 1):
        values = [_json_value(value) for value in frame.iloc[:, index - 1].dropna().head(5)]
        columns.append(
            TemplateColumn(
                name=name,
                column_index=index,
                inferred_type=_infer_values_type(list(frame.iloc[:, index - 1].dropna().head(5))),
                example_values=values,
            )
        )
    fingerprint = _template_fingerprint("csv", None, 1, headers, columns, [])
    return TemplateInspection(
        filename=Path(filename).name,
        file_type="csv",
        target_sheet=None,
        header_row=1,
        headers=headers,
        columns=columns,
        sheets=[],
        formula_cell_count=0,
        has_macros=False,
        fingerprint=fingerprint,
    )


def _validate_workbook_archive(content: bytes) -> None:
    try:
        validate_workbook_archive(content)
    except (WorkbookArchiveTooLargeError, WorkbookArchiveError) as exc:
        message = str(exc)
        if message == "The XLSX file is not a valid workbook archive.":
            message = "The target XLSX file is not a valid workbook archive."
        raise TemplateInspectionError(message) from exc


def _detect_header_row(worksheet) -> int:
    candidates = []
    for row_index in range(1, min(worksheet.max_row, 25) + 1):
        values = [worksheet.cell(row_index, column).value for column in range(1, worksheet.max_column + 1)]
        nonempty = [value for value in values if value is not None and str(value).strip()]
        if not nonempty:
            continue
        strings = sum(isinstance(value, str) and not value.startswith("=") for value in nonempty)
        unique = len({str(value).strip().casefold() for value in nonempty})
        candidates.append((strings * 3 + unique + len(nonempty), -row_index, row_index))
    if not candidates:
        raise TemplateInspectionError("Could not find a non-empty header row in the target worksheet.")
    return max(candidates)[2]


def _headers_from_row(worksheet, row_index: int) -> list[str]:
    last = max((cell.column for cell in worksheet[row_index] if cell.value is not None and str(cell.value).strip()), default=0)
    if not last:
        raise TemplateInspectionError("The selected target header row is empty.")
    headers = []
    for column in range(1, last + 1):
        value = worksheet.cell(row_index, column).value
        if value is None or not str(value).strip():
            raise TemplateInspectionError("Target headers must be contiguous and non-empty.")
        headers.append(str(value).strip())
    _validate_headers(headers)
    return headers


def _validate_headers(headers: list[str]) -> None:
    if not headers or len(headers) > MAX_TEMPLATE_COLUMNS:
        raise TemplateInspectionError("The target template must contain between 1 and 500 columns.")
    normalized = [name.casefold() for name in headers]
    if len(normalized) != len(set(normalized)):
        raise TemplateInspectionError("Target headers must be unique (case-insensitive).")
    if any(name.startswith("Unnamed:") for name in headers):
        raise TemplateInspectionError("Target headers must be non-empty.")


def _inspect_template_column(worksheet, header_row: int, index: int, name: str) -> TemplateColumn:
    cells = [worksheet.cell(row, index) for row in range(header_row + 1, min(worksheet.max_row, header_row + 20) + 1)]
    formula_cell = next((cell for cell in cells if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("="))), None)
    examples = [_json_value(cell.value) for cell in cells if cell.value is not None and cell is not formula_cell][:5]
    inferred = "formula" if formula_cell is not None else _infer_values_type([cell.value for cell in cells if cell.value is not None])
    sample_cell = next((cell for cell in cells if cell.value is not None), worksheet.cell(header_row + 1, index))
    return TemplateColumn(
        name=name,
        column_index=index,
        inferred_type=inferred,
        number_format=sample_cell.number_format or None,
        example_values=examples,
        formula=str(formula_cell.value) if formula_cell is not None else None,
        formula_row=formula_cell.row if formula_cell is not None else None,
    )


def _validate_template_formula(formula: str) -> None:
    normalized = formula.casefold().replace(" ", "")
    blocked = ("webservice(", "hyperlink(", "dde", "cmd|", "powershell")
    if any(token in normalized for token in blocked):
        raise TemplateInspectionError("The target workbook contains an unsupported external or active formula.")
    if re.search(r"\[(?:\d+|[^\]]+\.(?:xlsx?|xlsm|xlsb|csv))\]", normalized):
        raise TemplateInspectionError("The target workbook contains an unsupported external or active formula.")


def _infer_values_type(values: list[Any]) -> str:
    meaningful = [value for value in values if value is not None]
    if not meaningful:
        return "unknown"
    if all(isinstance(value, bool) for value in meaningful):
        return "boolean"
    if all(isinstance(value, (datetime, date)) for value in meaningful):
        return "date"
    if all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in meaningful):
        return "number"
    return "text"


def _template_fingerprint(file_type, sheet, header_row, headers, columns, sheets) -> str:
    structural = {
        "file_type": file_type,
        "target_sheet": sheet,
        "header_row": header_row,
        "headers": headers,
        "columns": [
            {
                "name": item.name,
                "type": item.inferred_type,
                "number_format": item.number_format,
                "formula": item.formula,
                "formula_row": item.formula_row,
            }
            for item in columns
        ],
        "sheets": [{"name": item.name, "state": item.state} for item in sheets],
    }
    return sha256(json.dumps(structural, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def _inspect_source(source: SourcePayload) -> SourceInspection:
    dataset = load_dataset(source.filename, source.content(), source.sheet)
    inspection = inspect_dataset(dataset)
    candidate_keys = []
    if len(dataset.frame):
        for column in inspection.columns:
            series = dataset.frame[column]
            if not series.isna().any() and int(series.nunique(dropna=False)) == len(series):
                candidate_keys.append([column])
    fingerprint = sha256(
        json.dumps(
            {"columns": inspection.columns, "types": [item.data_type for item in inspection.column_details]},
            separators=(",", ":"),
        ).encode()
    ).hexdigest()
    return SourceInspection(role=source.role, inspection=inspection, candidate_keys=candidate_keys, fingerprint=fingerprint)


def _validate_unique_roles(sources: list[SourcePayload]) -> None:
    roles = [source.role for source in sources]
    if len(roles) != len(set(roles)):
        raise TemplateInspectionError("Source roles must be unique.")


def _propose_mapping(column, sources, inspections, explicit, aliases) -> FieldMapping:
    if column.formula is not None:
        return FieldMapping(
            target_field=column.name,
            mapping_type="template_formula",
            confidence=1,
            evidence="Trusted formula found in the target template.",
            status="confirmed",
        )
    if explicit is not None:
        try:
            role, field = explicit.split(".", 1)
        except ValueError:
            return _unresolved(column.name, "incompatible", "Explicit mappings must use role.field syntax.")
        source = sources.get(role)
        inspection = inspections.get(role)
        if source is None or inspection is None or field not in inspection.inspection.columns:
            return _unresolved(column.name, "incompatible", f"Explicit source '{explicit}' does not exist.")
        if not _type_compatible(column.inferred_type, _source_dtype(inspections[role], field)):
            return FieldMapping(
                target_field=column.name, source_role=role, source_field=field, mapping_type="explicit", confidence=0,
                evidence=f"Explicit source '{explicit}' is incompatible with target type {column.inferred_type}.", status="incompatible"
            )
        return FieldMapping(
            target_field=column.name, source_role=role, source_field=field, mapping_type="explicit", confidence=1,
            evidence="Explicit user-confirmed mapping.", status="confirmed"
        )

    target_normalized = _normalize_name(column.name)
    candidates = _all_source_columns(inspections)
    exact = [item for item in candidates if item[2].casefold() == column.name.casefold()]
    if len(exact) == 1:
        return _candidate_mapping(column, exact[0], "exact", 0.99)
    normalized = [item for item in candidates if _normalize_name(item[2]) == target_normalized]
    if len(normalized) == 1:
        return _candidate_mapping(column, normalized[0], "normalized", 0.97)
    if len(exact) > 1 or len(normalized) > 1:
        options = exact or normalized
        return _ambiguous(column.name, options, "Multiple source fields have the same normalized target name.")

    accepted_aliases = {_normalize_name(value) for value in aliases}
    accepted_aliases.update(DOCUMENTED_ALIASES.get(target_normalized, set()))
    alias_matches = [item for item in candidates if _normalize_name(item[2]) in accepted_aliases]
    if len(alias_matches) == 1:
        return _candidate_mapping(column, alias_matches[0], "documented_alias", 0.94)
    if len(alias_matches) > 1:
        return _ambiguous(column.name, alias_matches, "Multiple documented aliases are present; business meaning must be confirmed.")

    target_tokens = _name_tokens(column.name)
    semantic = [
        item for item in candidates
        if target_tokens and target_tokens.intersection(_name_tokens(item[2])) and _type_compatible(column.inferred_type, item[3])
    ]
    if len(semantic) == 1:
        return _candidate_mapping(column, semantic[0], "semantic_candidate", 0.75, status="ambiguous")
    if semantic:
        return _ambiguous(column.name, semantic, "Several type-compatible semantic candidates overlap the target name.")
    return _unresolved(column.name, "missing", "No defensible source field was found.")


def _all_source_columns(inspections):
    return [
        (role, item.inspection.filename, detail.name, detail.data_type)
        for role, item in inspections.items()
        for detail in item.inspection.column_details
    ]


def _candidate_mapping(column, candidate, mapping_type, confidence, status="proposed_high_confidence"):
    role, filename, field, dtype = candidate
    if not _type_compatible(column.inferred_type, dtype):
        return FieldMapping(
            target_field=column.name, source_role=role, source_field=field, mapping_type=mapping_type,
            confidence=0, evidence=f"{filename}:{field} is incompatible with target type {column.inferred_type}.", status="incompatible"
        )
    return FieldMapping(
        target_field=column.name, source_role=role, source_field=field, mapping_type=mapping_type,
        confidence=confidence, evidence=f"Matched {filename}:{field} by {mapping_type.replace('_', ' ')}.", status=status
    )


def _ambiguous(target, candidates, reason):
    options = ", ".join(f"{item[0]}.{item[2]}" for item in candidates)
    return FieldMapping(
        target_field=target, mapping_type="unmapped", confidence=0, evidence=f"{reason} Candidates: {options}", status="ambiguous"
    )


def _unresolved(target, status, reason):
    return FieldMapping(target_field=target, mapping_type="unmapped", confidence=0, evidence=reason, status=status)


def _normalize_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.casefold())


def _name_tokens(value: str) -> set[str]:
    expanded = re.sub(r"([a-z])([A-Z])", r"\1 \2", value)
    return {token for token in re.split(r"[^A-Za-z0-9]+", expanded.casefold()) if len(token) > 2}


def _source_dtype(inspection: SourceInspection, field: str) -> str:
    return next(item.data_type for item in inspection.inspection.column_details if item.name == field)


def _type_compatible(target_type: str, source_dtype: str) -> bool:
    if target_type in {"unknown", "text"}:
        return True
    lowered = source_dtype.casefold()
    if target_type == "number":
        return any(item in lowered for item in ("int", "float", "decimal")) or "object" in lowered or "string" in lowered
    if target_type == "date":
        return "date" in lowered or "object" in lowered or "string" in lowered
    if target_type == "boolean":
        return "bool" in lowered or "object" in lowered or "string" in lowered
    return True


def _clarifications(mappings, required_fields, derivation_targets):
    required = set(required_fields) - derivation_targets
    clarifications = []
    for mapping in mappings:
        if mapping.target_field not in required or mapping.status in {"confirmed", "proposed_high_confidence"}:
            continue
        code = {
            "ambiguous": "ambiguous_mapping",
            "missing": "missing_required_field",
            "incompatible": "incompatible_type",
        }.get(mapping.status, "ambiguous_mapping")
        options = re.findall(r"[A-Za-z][A-Za-z0-9_-]*\.[^,]+", mapping.evidence.split("Candidates:")[-1]) if "Candidates:" in mapping.evidence else []
        clarifications.append(ClarificationRequirement(code=code, target_field=mapping.target_field, candidate_options=options, reason=mapping.evidence))
    return clarifications


def _validate_plan_targets(plan):
    headers = set(plan.target_headers)
    mapping_targets = [item.target_field for item in plan.mappings]
    if len(mapping_targets) != len(set(mapping_targets)) or set(mapping_targets) != headers:
        raise TransformValidationError("Mappings must contain each target field exactly once.")
    derivation_targets = [item.target_field for item in plan.derivations]
    if len(derivation_targets) != len(set(derivation_targets)) or not set(derivation_targets).issubset(headers):
        raise TransformValidationError("Derivations must target unique fields in the target template.")
    if not set(plan.required_fields).issubset(headers) or not set(plan.unique_fields).issubset(headers):
        raise TransformValidationError("Required and unique fields must exist in the target template.")


def _join_sources(plan, loaded):
    first_role = plan.source_roles[0]
    merged = loaded[first_role].rename(columns=lambda value: f"{first_role}.{value}")
    used = {first_role}
    diagnostics = []
    for rule in plan.joins:
        if rule.right_role in used or rule.right_role not in loaded:
            raise TransformValidationError(f"Join right role '{rule.right_role}' is missing or already joined.")
        left_keys = [_resolve_join_key(merged.columns, used, key) for key in rule.left_on]
        right_keys = [f"{rule.right_role}.{key}" for key in rule.right_on]
        right = loaded[rule.right_role].rename(columns=lambda value: f"{rule.right_role}.{value}")
        try:
            joined = join_datasets(merged, right, JoinSpec(left_on=left_keys, right_on=right_keys, how=rule.how))
        except TransformationError as exc:
            raise TransformValidationError(str(exc)) from exc
        diagnostic = joined.diagnostics
        unmatched_percentage = diagnostic.left_unmatched_rows / max(diagnostic.left_input_rows, 1) * 100
        if rule.block_many_to_many and diagnostic.many_to_many_detected:
            raise TransformValidationError(f"Unsafe join to {rule.right_role}: many-to-many keys would multiply rows.")
        if rule.block_row_multiplication and diagnostic.row_multiplication_occurred:
            raise TransformValidationError(f"Unsafe join to {rule.right_role}: repeated keys would multiply rows.")
        if unmatched_percentage > rule.max_unmatched_left_percentage:
            raise TransformValidationError(
                f"Unsafe join to {rule.right_role}: {unmatched_percentage:.2f}% unmatched left rows exceeds the allowed {rule.max_unmatched_left_percentage:.2f}%."
            )
        merged = joined.frame
        diagnostics.append(diagnostic)
        used.add(rule.right_role)
    missing = [role for role in plan.source_roles if role not in used]
    if missing:
        raise TransformValidationError(f"Source roles are not connected by the join plan: {missing}.")
    return merged.reset_index(drop=True), diagnostics


def _resolve_join_key(columns, used_roles, key):
    if "." in key:
        if key not in columns:
            raise TransformValidationError(f"Join field '{key}' does not exist in the joined sources.")
        return key
    matches = [f"{role}.{key}" for role in used_roles if f"{role}.{key}" in columns]
    if len(matches) != 1:
        raise TransformValidationError(f"Join field '{key}' is missing or ambiguous; use role.field syntax.")
    return matches[0]


def _source_series(frame, role, field):
    name = f"{role}.{field}"
    if name not in frame.columns:
        raise TransformValidationError(f"Mapped source field '{name}' does not exist after joins.")
    return frame[name].copy()


def _apply_mapping_transformation(series, operation):
    if operation == "none":
        return series
    if operation == "trim":
        return series.map(lambda value: value.strip() if isinstance(value, str) else value)
    if operation == "upper":
        return series.map(lambda value: value.upper() if isinstance(value, str) else value)
    if operation == "lower":
        return series.map(lambda value: value.lower() if isinstance(value, str) else value)
    if operation in {"normalize_number", "normalize_currency"}:
        cleaned = series.astype("string")
        if operation == "normalize_currency":
            cleaned = cleaned.str.replace(r"[$,()]", lambda match: "-" if match.group(0) == "(" else "", regex=True).str.rstrip(")")
        else:
            cleaned = cleaned.str.replace(",", "", regex=False)
        return pd.to_numeric(cleaned, errors="coerce")
    if operation == "normalize_date":
        return pd.to_datetime(series, errors="coerce")
    raise TransformValidationError(f"Unsupported mapping transformation '{operation}'.")


def _coerce_target_type(series, target_type, target):
    if target_type in {"unknown", "text", "formula"}:
        return series
    original_non_null = int(series.notna().sum())
    if target_type == "number":
        converted = pd.to_numeric(series, errors="coerce")
    elif target_type == "date":
        converted = pd.to_datetime(series, errors="coerce")
    elif target_type == "boolean":
        normalized = series.map(lambda value: str(value).strip().casefold() if value is not None else None)
        converted = normalized.map({"true": True, "yes": True, "1": True, "false": False, "no": False, "0": False})
    else:
        return series
    if int(converted.notna().sum()) != original_non_null:
        raise TransformValidationError(f"Values for target field '{target}' are incompatible with expected type {target_type}.")
    return converted


def _derive(rule: DerivationRule, frame: pd.DataFrame, output: pd.DataFrame, evidence):
    inputs = []
    for item in rule.inputs:
        if item.source_role == "target":
            if item.source_field not in output.columns or output[item.source_field].isna().all():
                raise TransformValidationError(
                    f"Derivation '{rule.target_field}' references target field '{item.source_field}' before it is available."
                )
            source = output[item.source_field].copy()
        else:
            source = _source_series(frame, item.source_role, item.source_field)
        inputs.append(_apply_mapping_transformation(source, item.transformation))
    if rule.operation == "concat":
        result = inputs[0].fillna("").astype(str)
        for item in inputs[1:]:
            result = result.str.cat(item.fillna("").astype(str), sep=rule.separator)
        return result.str.strip(), []
    left = pd.to_numeric(inputs[0], errors="coerce")
    if int(left.notna().sum()) != int(inputs[0].notna().sum()):
        raise TransformValidationError(f"Derivation '{rule.target_field}' requires numeric input.")
    if rule.operation == "policy_multiply":
        rate, ids = _policy_rate(evidence, rule.policy_query or "")
        return left * rate, ids
    right = pd.to_numeric(inputs[1], errors="coerce") if len(inputs) == 2 else rule.constant
    if isinstance(right, pd.Series) and int(right.notna().sum()) != int(inputs[1].notna().sum()):
        raise TransformValidationError(f"Derivation '{rule.target_field}' requires numeric input.")
    if rule.operation == "add":
        return left + right, []
    if rule.operation == "subtract":
        return left - right, []
    if rule.operation == "multiply":
        return left * right, []
    if rule.operation == "divide":
        if (right == 0).any() if isinstance(right, pd.Series) else right == 0:
            raise TransformValidationError(f"Derivation '{rule.target_field}' would divide by zero.")
        return left / right, []
    raise TransformValidationError(f"Unsupported derivation '{rule.operation}'.")


def _policy_rate(evidence, query):
    if not evidence:
        raise TransformValidationError("A policy-grounded derivation requires retrieved policy evidence.")
    query_tokens = _name_tokens(query)
    relevant = [item for item in evidence if not query_tokens or query_tokens.intersection(_name_tokens(item.text))]
    if not relevant:
        raise TransformValidationError("Retrieved policy evidence does not support the requested derivation.")
    rates = []
    for item in relevant:
        rates.extend(float(match) / 100 for match in re.findall(r"(?<![\d.])(\d+(?:\.\d+)?)\s*%", item.text))
    unique = sorted(set(rates))
    if len(unique) != 1 or not 0 <= unique[0] <= 1:
        raise TransformValidationError("Policy evidence must establish exactly one unambiguous percentage rate.")
    return unique[0], [str(item.source.chunk_id) for item in relevant]


def _validate_output_frame(frame, plan, diagnostics):
    checks = [
        ValidationCheck(name="exact_schema_order", passed=list(frame.columns) == plan.target_headers, detail="Output columns match the target schema and order."),
        ValidationCheck(name="row_count", passed=len(frame) >= 0, detail=f"Output contains {len(frame)} rows."),
    ]
    errors = []
    warnings = []
    formula_fields = {mapping.target_field for mapping in plan.mappings if mapping.mapping_type == "template_formula"}
    for field in plan.required_fields:
        if field in formula_fields:
            missing = 0
        else:
            values = frame[field]
            missing = int((values.isna() | values.map(lambda value: isinstance(value, str) and not value.strip())).sum())
        passed = missing == 0
        checks.append(ValidationCheck(name=f"required:{field}", passed=passed, detail=f"{missing} missing required values in {field}."))
        if not passed:
            errors.append(f"Required field '{field}' contains {missing} missing values.")
    for field in plan.unique_fields:
        duplicates = int(frame[field].duplicated(keep=False).sum())
        passed = duplicates == 0
        checks.append(ValidationCheck(name=f"unique:{field}", passed=passed, detail=f"{duplicates} rows violate uniqueness for {field}."))
        if not passed:
            errors.append(f"Unique field '{field}' contains duplicate values.")
    for index, diagnostic in enumerate(diagnostics, 1):
        if diagnostic.left_unmatched_rows or diagnostic.right_unmatched_rows:
            warnings.append(
                f"Join {index} left {diagnostic.left_unmatched_rows} source rows and {diagnostic.right_unmatched_rows} lookup rows unmatched."
            )
    return errors, warnings, checks


def _write_artifact(request, template, frame, provenance):
    safe = frame.copy()
    for column in safe.columns:
        safe[column] = safe[column].map(_safe_cell)
    output = BytesIO()
    extension = Path(request.target.filename).suffix.lower()
    if extension == ".csv":
        content = safe.to_csv(index=False).encode("utf-8")
        media_type = "text/csv"
        format_name = "csv"
    else:
        workbook = load_workbook(BytesIO(request.target.content()), data_only=False, read_only=False, keep_links=False)
        worksheet = workbook[template.target_sheet]
        data_row = template.header_row + 1
        style_row = data_row
        formula_by_column = {
            item.column_index: (item.formula, item.formula_row or style_row)
            for item in template.columns
            if item.formula
        }
        style_cells = {index: copy(worksheet.cell(style_row, index)._style) for index in range(1, len(template.headers) + 1)}
        formats = {index: worksheet.cell(style_row, index).number_format for index in range(1, len(template.headers) + 1)}
        for row in range(data_row, worksheet.max_row + 1):
            for column in range(1, len(template.headers) + 1):
                worksheet.cell(row, column).value = None
        for offset, (_, record) in enumerate(safe.iterrows()):
            row = data_row + offset
            for column, header in enumerate(template.headers, 1):
                cell = worksheet.cell(row, column)
                cell._style = copy(style_cells[column])
                cell.number_format = formats[column]
                formula_template = formula_by_column.get(column)
                if formula_template:
                    formula, formula_row = formula_template
                    try:
                        cell.value = Translator(
                            formula,
                            origin=worksheet.cell(formula_row, column).coordinate,
                        ).translate_formula(cell.coordinate)
                    except Exception as exc:
                        raise TransformValidationError(f"Could not safely translate trusted template formula for '{header}'.") from exc
                else:
                    cell.value = _excel_value(record[header])
        workbook.save(output)
        content = output.getvalue()
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        format_name = "xlsx"
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", Path(request.target.filename).stem).strip("_") or "submission"
    return GeneratedArtifact(
        filename=f"{stem}_completed.{format_name}",
        format=format_name,
        media_type=media_type,
        row_count=len(frame),
        column_count=len(frame.columns),
        content=content,
        artifact_id=uuid4(),
        created_at=datetime.now(timezone.utc),
        verification_status="completed",
        provenance={item.target_field: f"{', '.join(item.source_fields) or 'template'} -> {item.transformation}" for item in provenance},
    )


def _validate_artifact_reopens(artifact, template, expected_rows):
    checks = []
    try:
        if artifact.format == "csv":
            reopened = pd.read_csv(BytesIO(artifact.content), encoding="utf-8-sig")
            schema_ok = list(reopened.columns) == template.headers
            rows_ok = len(reopened) == expected_rows
            checks.extend([
                ValidationCheck(name="artifact_reopens", passed=True, detail="CSV artifact reopened successfully."),
                ValidationCheck(name="artifact_schema", passed=schema_ok, detail="Reopened CSV schema/order matches the target."),
                ValidationCheck(name="artifact_rows", passed=rows_ok, detail=f"Reopened CSV contains {len(reopened)} rows."),
            ])
        else:
            reopened = load_workbook(BytesIO(artifact.content), data_only=False, read_only=False, keep_links=False)
            sheets_ok = reopened.sheetnames == [item.name for item in template.sheets]
            worksheet = reopened[template.target_sheet]
            headers_ok = [worksheet.cell(template.header_row, index).value for index in range(1, len(template.headers) + 1)] == template.headers
            formula_columns = [item for item in template.columns if item.formula]
            formulas_ok = all(
                all(
                    worksheet.cell(template.header_row + 1 + offset, item.column_index).value
                    == Translator(
                        item.formula,
                        origin=worksheet.cell(item.formula_row or template.header_row + 1, item.column_index).coordinate,
                    ).translate_formula(worksheet.cell(template.header_row + 1 + offset, item.column_index).coordinate)
                    for offset in range(expected_rows)
                )
                for item in formula_columns
            )
            first_trailing_row = template.header_row + expected_rows + 1
            stale_rows_cleared = all(
                worksheet.cell(row, column).value is None
                for row in range(first_trailing_row, worksheet.max_row + 1)
                for column in range(1, len(template.headers) + 1)
            )
            checks.extend([
                ValidationCheck(name="artifact_reopens", passed=True, detail="XLSX artifact reopened successfully."),
                ValidationCheck(name="workbook_sheets_preserved", passed=sheets_ok, detail="Workbook sheet names and order were preserved."),
                ValidationCheck(name="artifact_schema", passed=headers_ok, detail="Reopened XLSX schema/order matches the target."),
                ValidationCheck(name="artifact_rows", passed=stale_rows_cleared, detail="No stale target rows remain after the emitted data region."),
                ValidationCheck(name="trusted_formulas_preserved", passed=formulas_ok, detail="Trusted target formulas were preserved for output rows."),
            ])
    except Exception:
        checks.append(ValidationCheck(name="artifact_reopens", passed=False, detail="Generated artifact could not be reopened safely."))
    return checks


def _safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(FORMULA_PREFIXES):
        return f"'{value}"
    return value


def _excel_value(value):
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item"):
        return value.item()
    return value


def _json_value(value):
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if hasattr(value, "item"):
        return value.item()
    return value if isinstance(value, (str, int, float, bool)) else str(value)
