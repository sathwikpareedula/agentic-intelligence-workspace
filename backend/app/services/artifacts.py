"""In-memory tabular artifact generation."""

from dataclasses import dataclass, replace
from io import BytesIO
from pathlib import Path
import re
from datetime import datetime, timezone
from typing import Protocol
from uuid import UUID, uuid4

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


@dataclass(frozen=True)
class GeneratedArtifact:
    filename: str
    format: str
    media_type: str
    row_count: int
    column_count: int
    content: bytes
    artifact_id: UUID
    created_at: datetime
    producing_task_id: UUID | None = None
    producing_workflow_id: UUID | None = None
    producing_run_id: UUID | None = None
    verification_status: str = "not_verified"
    provenance: dict[str, str] | None = None


def generate_artifact(frame: pd.DataFrame, source_filename: str, format: str) -> GeneratedArtifact:
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", Path(source_filename).stem).strip("_") or "dataset"
    filename = f"{stem}_result.{format}"
    safe_frame = _sanitize_spreadsheet_strings(frame)
    output = BytesIO()
    if format == "csv":
        content = safe_frame.to_csv(index=False).encode("utf-8")
        media_type = "text/csv"
    elif format == "xlsx":
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            safe_frame.to_excel(writer, sheet_name="Result", index=False)
        content = output.getvalue()
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        raise ValueError("Artifact format must be csv or xlsx.")
    return GeneratedArtifact(filename, format, media_type, len(frame), len(frame.columns), content, uuid4(), datetime.now(timezone.utc))


def generate_management_workbook(
    frame: pd.DataFrame,
    source_filename: str,
    *,
    task_id: UUID | None = None,
    workflow_id: UUID | None = None,
    provenance: dict[str, str] | None = None,
) -> GeneratedArtifact:
    """Produce a deterministic management workbook with data, summary, and provenance."""
    safe_frame = _sanitize_spreadsheet_strings(frame)
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        safe_frame.to_excel(writer, sheet_name="Data", index=False)
        summary = pd.DataFrame(
            {"Metric": ["Rows", "Columns"], "Value": [len(frame), len(frame.columns)]}
        )
        summary.to_excel(writer, sheet_name="Summary", index=False)
        provenance_frame = pd.DataFrame(
            list((provenance or {}).items()), columns=["Property", "Value"]
        )
        provenance_frame.to_excel(writer, sheet_name="Provenance", index=False)
        numeric_columns = [index + 1 for index, column in enumerate(frame.columns) if pd.api.types.is_numeric_dtype(frame[column])]
        if numeric_columns and len(frame) > 0:
            chart = BarChart()
            first_numeric = numeric_columns[0]
            chart.add_data(Reference(writer.book["Data"], min_col=first_numeric, min_row=1, max_row=len(frame) + 1), titles_from_data=True)
            chart.title = f"{frame.columns[first_numeric - 1]} overview"
            writer.book["Summary"].add_chart(chart, "D2")
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", Path(source_filename).stem).strip("_") or "report"
    return GeneratedArtifact(
        filename=f"{stem}_management_report.xlsx",
        format="xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        row_count=len(frame),
        column_count=len(frame.columns),
        content=output.getvalue(),
        artifact_id=uuid4(),
        created_at=datetime.now(timezone.utc),
        producing_task_id=task_id,
        producing_workflow_id=workflow_id,
        provenance=provenance,
    )


def generate_sales_management_workbook(
    cleaned_transactions: pd.DataFrame,
    regional_performance: pd.DataFrame,
    commissions: pd.DataFrame,
    *,
    reporting_period: str,
    data_quality: dict[str, object],
    customer_join_diagnostics,
    target_join_diagnostics: dict[str, int],
    citations: list,
    provenance: dict[str, str],
    warnings: list[str],
) -> GeneratedArtifact:
    """Produce a polished, auditable management workbook from deterministic tables."""
    output = BytesIO()
    targeted = regional_performance[regional_performance["target"].notna()]
    largest = targeted.sort_values(
        ["underperformance", "region"], ascending=[False, True], kind="stable"
    ).iloc[0] if not targeted.empty else None
    summary = pd.DataFrame(
        {
            "Metric": [
                "Report scope",
                f"Completed {reporting_period} transactions",
                "Source transaction rows",
                "Regions represented",
                "Salespeople represented",
                "Total net sales",
                "Total regional target",
                "Net target variance",
                "Largest underperforming region",
                "Largest regional shortfall",
                "Total commission",
                "Verification note",
            ],
            "Value": [
                f"Completed {reporting_period} net sales after discounts",
                len(cleaned_transactions),
                data_quality.get("original_transaction_rows", len(cleaned_transactions)),
                len(regional_performance),
                len(commissions),
                float(cleaned_transactions["net_sales"].sum()),
                float(targeted["target"].sum()),
                float(targeted["variance"].sum()),
                largest["region"] if largest is not None else "Not available",
                float(largest["underperformance"]) if largest is not None else "Not available",
                float(commissions["commission"].sum()),
                "Numeric values are deterministic; commissions require the cited policy evidence.",
            ],
        }
    )
    diagnostic_rows = [
        {"Category": "Data quality", "Metric": key.replace("_", " ").title(), "Value": value, "Severity": "Info"}
        for key, value in data_quality.items()
    ]
    diagnostic_rows.extend(
        {"Category": "Customer join", "Metric": key.replace("_", " ").title(), "Value": value, "Severity": "Info"}
        for key, value in customer_join_diagnostics.model_dump(mode="json").items()
        if key != "warning"
    )
    diagnostic_rows.extend(
        {"Category": "Target join", "Metric": key.replace("_", " ").title(), "Value": value, "Severity": "Info"}
        for key, value in target_join_diagnostics.items()
    )
    diagnostic_rows.extend(
        {"Category": "Warning", "Metric": "Attention required", "Value": warning, "Severity": "Warning"}
        for warning in warnings
    )
    diagnostics = pd.DataFrame(diagnostic_rows)
    provenance_rows = [
        {"Source / property": key.replace("_", " ").title(), "Value": value, "Page": None, "Chunk ID": None}
        for key, value in provenance.items()
    ]
    provenance_rows.extend(
        {
            "Source / property": source.filename,
            "Value": "Commission policy evidence used by the deterministic calculator",
            "Page": source.page_number,
            "Chunk ID": str(source.chunk_id),
        }
        for source in citations
    )
    provenance_frame = pd.DataFrame(provenance_rows)
    frames = {
        "Executive Summary": _sanitize_spreadsheet_strings(summary),
        "Regional Performance": _sanitize_spreadsheet_strings(regional_performance),
        "Salesperson Performance": _sanitize_spreadsheet_strings(commissions),
        "Cleaned Transactions": _sanitize_spreadsheet_strings(cleaned_transactions),
        "Data Quality": _sanitize_spreadsheet_strings(diagnostics),
        "Provenance & Sources": _sanitize_spreadsheet_strings(provenance_frame),
    }
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, frame in frames.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
        regional_sheet = writer.book["Regional Performance"]
        header_index = {cell.value: cell.column for cell in regional_sheet[1]}
        comparison_chart = BarChart()
        comparison_chart.type = "col"
        comparison_chart.style = 10
        comparison_chart.title = "Actual sales vs target by region"
        comparison_chart.y_axis.title = "Sales amount"
        comparison_chart.x_axis.title = "Region"
        comparison_chart.height = 8
        comparison_chart.width = 15
        comparison_chart.add_data(
            Reference(
                regional_sheet,
                min_col=header_index["net_sales"],
                max_col=header_index["target"],
                min_row=1,
                max_row=len(regional_performance) + 1,
            ),
            titles_from_data=True,
        )
        comparison_chart.set_categories(
            Reference(regional_sheet, min_col=header_index["region"], min_row=2, max_row=len(regional_performance) + 1)
        )
        writer.book["Executive Summary"].add_chart(comparison_chart, "D2")

        shortfall_chart = BarChart()
        shortfall_chart.type = "bar"
        shortfall_chart.style = 12
        shortfall_chart.title = "Regional shortfall"
        shortfall_chart.x_axis.title = "Shortfall amount"
        shortfall_chart.height = 7
        shortfall_chart.width = 12
        shortfall_chart.add_data(
            Reference(regional_sheet, min_col=header_index["underperformance"], min_row=1, max_row=len(regional_performance) + 1),
            titles_from_data=True,
        )
        shortfall_chart.set_categories(
            Reference(regional_sheet, min_col=header_index["region"], min_row=2, max_row=len(regional_performance) + 1)
        )
        regional_sheet.add_chart(shortfall_chart, "K2")

        salesperson_sheet = writer.book["Salesperson Performance"]
        salesperson_headers = {cell.value: cell.column for cell in salesperson_sheet[1]}
        commission_chart = BarChart()
        commission_chart.type = "col"
        commission_chart.style = 11
        commission_chart.title = "Commission by salesperson"
        commission_chart.y_axis.title = "Commission"
        commission_chart.height = 7
        commission_chart.width = 12
        commission_chart.add_data(
            Reference(salesperson_sheet, min_col=salesperson_headers["commission"], min_row=1, max_row=len(commissions) + 1),
            titles_from_data=True,
        )
        commission_chart.set_categories(
            Reference(salesperson_sheet, min_col=salesperson_headers["salesperson"], min_row=2, max_row=len(commissions) + 1)
        )
        salesperson_sheet.add_chart(commission_chart, "G2")

        _style_sales_workbook(writer.book, frames)

    return GeneratedArtifact(
        filename=f"{reporting_period.split()[0].lower()}_sales_management_report.xlsx",
        format="xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        row_count=len(regional_performance),
        column_count=len(regional_performance.columns),
        content=output.getvalue(),
        artifact_id=uuid4(),
        created_at=datetime.now(timezone.utc),
        provenance=provenance,
    )


def _style_sales_workbook(workbook, frames: dict[str, pd.DataFrame]) -> None:
    header_fill = PatternFill("solid", fgColor="174C3C")
    header_font = Font(color="FFFFFF", bold=True)
    warning_fill = PatternFill("solid", fgColor="FFF1D6")
    for sheet_name, frame in frames.items():
        sheet = workbook[sheet_name]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 24
        for index, column in enumerate(frame.columns, 1):
            values = [str(column)] + [str(value) for value in frame[column].dropna().head(100)]
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(max(len(value) for value in values) + 2, 46)
        if len(frame):
            table_name = re.sub(r"[^A-Za-z0-9]", "", sheet_name) + "Table"
            table = Table(displayName=table_name, ref=f"A1:{sheet.cell(len(frame) + 1, len(frame.columns)).coordinate}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
            sheet.add_table(table)

    # Long source/diagnostic text must remain readable in the exported workbook.
    for name, column in (("Executive Summary", "B"), ("Data Quality", "C"), ("Provenance & Sources", "B")):
        sheet = workbook[name]
        sheet.column_dimensions[column].width = 64
        for row in range(2, sheet.max_row + 1):
            cell = sheet[f"{column}{row}"]
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[row].height = max(22, 16 * ((len(str(cell.value or "")) // 60) + 1))

    for sheet_name in ("Regional Performance", "Salesperson Performance", "Cleaned Transactions"):
        sheet = workbook[sheet_name]
        headers = {cell.value: cell.column for cell in sheet[1]}
        for name in ("amount", "discount", "net_sales", "target", "variance", "underperformance", "commission"):
            if name in headers:
                for row in range(2, sheet.max_row + 1):
                    sheet.cell(row, headers[name]).number_format = '#,##0.00;[Red]-#,##0.00'
        for name in ("attainment_pct", "share_of_sales_pct", "share_of_shortfall_pct"):
            if name in headers:
                for row in range(2, sheet.max_row + 1):
                    sheet.cell(row, headers[name]).number_format = '0.0"%"'
    quality_sheet = workbook["Data Quality"]
    quality_headers = {cell.value: cell.column for cell in quality_sheet[1]}
    if "Severity" in quality_headers:
        for row in range(2, quality_sheet.max_row + 1):
            if quality_sheet.cell(row, quality_headers["Severity"]).value == "Warning":
                for cell in quality_sheet[row]:
                    cell.fill = warning_fill


class InMemoryArtifactRepository:
    """Request-independent test repository behind a small persistence boundary."""

    def __init__(self) -> None:
        self._artifacts: dict[UUID, GeneratedArtifact] = {}

    def save(self, artifact: GeneratedArtifact) -> None:
        self._artifacts[artifact.artifact_id] = artifact

    def get(self, artifact_id: UUID) -> GeneratedArtifact | None:
        return self._artifacts.get(artifact_id)

    def link_to_execution(self, artifact_ids: list[UUID], task_id: UUID, verification_status: str) -> None:
        for artifact_id in artifact_ids:
            artifact = self._artifacts.get(artifact_id)
            if artifact is not None:
                self._artifacts[artifact_id] = replace(
                    artifact,
                    producing_task_id=task_id,
                    producing_run_id=task_id,
                    verification_status=verification_status,
                )

    def link_to_workflow(self, artifact_ids: list[UUID], workflow_id: UUID, run_id: UUID) -> None:
        for artifact_id in artifact_ids:
            artifact = self._artifacts.get(artifact_id)
            if artifact is not None:
                self._artifacts[artifact_id] = replace(
                    artifact,
                    producing_workflow_id=workflow_id,
                    producing_run_id=run_id,
                )


class ArtifactRepository(Protocol):
    def save(self, artifact: GeneratedArtifact) -> None: ...

    def get(self, artifact_id: UUID) -> GeneratedArtifact | None: ...

    def link_to_execution(self, artifact_ids: list[UUID], task_id: UUID, verification_status: str) -> None: ...

    def link_to_workflow(self, artifact_ids: list[UUID], workflow_id: UUID, run_id: UUID) -> None: ...


def _sanitize_spreadsheet_strings(frame: pd.DataFrame) -> pd.DataFrame:
    """Prevent uploaded text from becoming a spreadsheet formula on export."""

    result = frame.copy()
    for column in result.columns:
        result[column] = result[column].map(_sanitize_cell)
    return result


def _sanitize_cell(value: object) -> object:
    if not isinstance(value, str):
        return value
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value
