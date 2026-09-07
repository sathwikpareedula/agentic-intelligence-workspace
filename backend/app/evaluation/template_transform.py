"""Controlled offline evaluation for transform-to-template trust behavior."""

from __future__ import annotations

import argparse
import base64
from io import BytesIO
import json
from pathlib import Path
from uuid import uuid4

import pandas as pd
from openpyxl import Workbook, load_workbook

from app.agent.tools import ToolRegistry, template_transform_tool
from app.models.grades import PolicyEvidence
from app.models.retrieval import SourceReference
from app.models.template_transforms import (
    DerivationRule,
    FieldReference,
    FilePayload,
    JoinRule,
    SourcePayload,
    TransformExecutionRequest,
    TransformProposalRequest,
)
from app.models.workflows import WorkflowCreate, WorkflowStep
from app.services.artifacts import InMemoryArtifactRepository
from app.services.template_transforms import TransformValidationError, execute_transform, propose_transform
from app.services.workflows import InMemoryWorkflowRepository, WorkflowService


def evaluate_template_transform_cases(path: Path) -> dict:
    specification = json.loads(path.read_text(encoding="utf-8"))
    request, proposal = _ready_request()
    executed = execute_transform(request)
    workbook = load_workbook(BytesIO(executed.artifact.content), data_only=False)
    sheet = workbook["Submission"]
    by_target = {item.target_field: item for item in proposal.plan.mappings}

    ambiguous = _ambiguous_proposal()
    missing = _missing_proposal()
    incompatible_refused = _incompatible_refused()
    suspicious_refused = _suspicious_join_refused()
    unsupported_policy_refused = _unsupported_policy_refused(request)

    artifacts = InMemoryArtifactRepository()
    workflows = WorkflowService(InMemoryWorkflowRepository(), ToolRegistry([template_transform_tool(artifacts)]), artifacts)
    workflow = workflows.create(
        WorkflowCreate(
            name="Monthly submission",
            steps=[WorkflowStep(tool="template.transform", arguments=request.model_dump(mode="json"))],
        )
    )
    rerun = workflows.rerun(workflow.workflow_id, {})
    drift_source = _source(
        "raw_orders.csv",
        "orders",
        _csv([{"order_id": 1, "customer_id": "C1", "gross_sales": "$10.00", "returns": 1, "unexpected": 1}]),
    )
    schema_drift = workflows.rerun(
        workflow.workflow_id,
        {1: {"sources": [drift_source.model_dump(mode="json"), request.sources[1].model_dump(mode="json")]}}
    )
    changed_target = FilePayload(
        filename="required_template.xlsx",
        content_base64=base64.b64encode(_template(extra_header=True)).decode(),
        sheet="Submission",
    )
    template_drift = workflows.rerun(workflow.workflow_id, {1: {"target": changed_target.model_dump(mode="json")}})

    observed = {
        "A_exact_mapping": _simple_mapping("order_id", "order_id") == "exact",
        "B_normalized_mapping": by_target["Order ID"].mapping_type == "normalized",
        "C_ambiguous_mapping": ambiguous.status == "clarification_required" and ambiguous.clarifications[0].code == "ambiguous_mapping",
        "D_missing_required_field": missing.status == "clarification_required" and missing.clarifications[0].code == "missing_required_field",
        "E_incompatible_type": incompatible_refused,
        "F_safe_join": executed.validation.join_diagnostics and not executed.validation.join_diagnostics[0].row_multiplication_occurred,
        "G_suspicious_join": suspicious_refused,
        "H_derived_field": sheet["D2"].value == 900 and sheet["E2"].value == 45,
        "I_exact_output_schema_order": [sheet.cell(1, index).value for index in range(1, 7)] == request.plan.target_headers,
        "J_workbook_reopen": any(item.name == "artifact_reopens" and item.passed for item in executed.validation.checks),
        "K_template_preservation": workbook.sheetnames == ["Submission", "Instructions"] and workbook["Instructions"]["A1"].value == "Preserve me",
        "L_formula_injection_protection": isinstance(sheet["B3"].value, str) and sheet["B3"].value.startswith("'="),
        "M_policy_grounded_derivation": bool(next(item for item in executed.provenance if item.target_field == "Commission").policy_evidence_ids),
        "N_unsupported_policy_refusal": unsupported_policy_refused,
        "O_workflow_rerun": rerun.status == "completed" and bool(rerun.observations[0].artifact_ids),
        "P_schema_drift": schema_drift.status == "failed" and bool(schema_drift.error and "Schema drift" in schema_drift.error),
        "Q_template_drift": template_drift.status == "failed" and bool(template_drift.error and "Template drift" in template_drift.error),
        "R_confirmed_mapping_reuse": rerun.status == "completed" and request.plan.mappings[0].status == "confirmed",
    }
    cases = []
    for expected in specification["cases"]:
        case_id = expected["id"]
        cases.append({"id": case_id, "name": expected["name"], "passed": bool(observed.get(case_id)), "observed": observed.get(case_id)})
    passed = sum(item["passed"] for item in cases)
    return {
        "case_count": len(cases),
        "passed_count": passed,
        "failed_count": len(cases) - passed,
        "case_pass_rate": passed / len(cases) if cases else 0,
        "cases": cases,
        "scope": "Controlled deterministic CSV/XLSX transform-to-template evaluation; no hosted model, hosted embeddings, or live database.",
    }


def _ready_request():
    target_content = _template()
    sources = [_orders(), _customers()]
    target = FilePayload(filename="required_template.xlsx", content_base64=base64.b64encode(target_content).decode(), sheet="Submission")
    proposal = propose_transform(TransformProposalRequest(target=target, sources=sources))
    mappings = [
        item if item.target_field in {"Net Sales", "Commission"} else item.model_copy(update={"status": "confirmed"})
        for item in proposal.plan.mappings
    ]
    plan = proposal.plan.model_copy(
        update={
            "mappings": mappings,
            "joins": [JoinRule(right_role="customers", left_on=["orders.customer_id"], right_on=["cust_id"])],
            "derivations": [
                DerivationRule(
                    target_field="Net Sales",
                    operation="subtract",
                    inputs=[
                        FieldReference(source_role="orders", source_field="gross_sales", transformation="normalize_currency"),
                        FieldReference(source_role="orders", source_field="returns"),
                    ],
                ),
                DerivationRule(
                    target_field="Commission",
                    operation="policy_multiply",
                    inputs=[FieldReference(source_role="target", source_field="Net Sales")],
                    policy_query="commission policy rate",
                ),
            ],
            "unique_fields": ["Order ID"],
        }
    )
    request = TransformExecutionRequest(target=target, sources=sources, plan=plan, policy_evidence=[_evidence()])
    return request, proposal


def _template(extra_header: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Submission"
    headers = ["Order ID", "Customer Name", "Region", "Net Sales", "Commission", "Total"]
    if extra_header:
        headers.append("Unexpected")
    sheet.append(headers)
    row = [1, "Example", "North", 100.0, 5.0, "=D2+E2"]
    if extra_header:
        row.append("value")
    sheet.append(row)
    sheet["D2"].number_format = "$#,##0.00"
    sheet["E2"].number_format = "$#,##0.00"
    workbook.create_sheet("Instructions")["A1"] = "Preserve me"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _csv(rows):
    return pd.DataFrame(rows).to_csv(index=False).encode()


def _source(filename, role, content):
    return SourcePayload(filename=filename, role=role, content_base64=base64.b64encode(content).decode())


def _orders():
    return _source(
        "raw_orders.csv",
        "orders",
        _csv([
            {"order_id": 101, "customer_id": "C1", "gross_sales": "$1,000.00", "returns": 100},
            {"order_id": 102, "customer_id": "C2", "gross_sales": "$500.00", "returns": 50},
        ]),
    )


def _customers(duplicate=False):
    rows = [
        {"cust_id": "C1", "customer_name": "Alice", "territory": "North"},
        {"cust_id": "C2", "customer_name": "=HYPERLINK(\"x\")", "territory": "South"},
    ]
    if duplicate:
        rows.append({"cust_id": "C1", "customer_name": "Alicia", "territory": "West"})
    return _source("customer_master.csv", "customers", _csv(rows))


def _evidence(text="The commission policy rate is 5% of eligible sales."):
    return PolicyEvidence(
        text=text,
        source=SourceReference(document_id=uuid4(), filename="reporting_policy.pdf", page_number=1, chunk_id=uuid4()),
    )


def _simple_mapping(target_name, source_name):
    target = FilePayload(filename="target.csv", content_base64=base64.b64encode(f"{target_name}\n".encode()).decode())
    source = _source("source.csv", "source", _csv([{source_name: 1}]))
    proposal = propose_transform(TransformProposalRequest(target=target, sources=[source]))
    return proposal.plan.mappings[0].mapping_type


def _ambiguous_proposal():
    target = FilePayload(filename="target.csv", content_base64=base64.b64encode(b"Region\n").decode())
    return propose_transform(
        TransformProposalRequest(
            target=target,
            sources=[
                _source("territories.csv", "territories", _csv([{"territory": "North"}])),
                _source("areas.csv", "areas", _csv([{"sales_area": "N"}])),
            ],
        )
    )


def _missing_proposal():
    target = FilePayload(filename="target.csv", content_base64=base64.b64encode(b"Manager Approval Code\n").decode())
    return propose_transform(TransformProposalRequest(target=target, sources=[_orders()]))


def _incompatible_refused():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Amount"])
    sheet.append([1.0])
    output = BytesIO()
    workbook.save(output)
    target = FilePayload(filename="target.xlsx", content_base64=base64.b64encode(output.getvalue()).decode())
    source = _source("source.csv", "source", _csv([{"amount": "not-a-number"}]))
    proposal = propose_transform(TransformProposalRequest(target=target, sources=[source]))
    mapping = proposal.plan.mappings[0].model_copy(update={"status": "confirmed"})
    try:
        execute_transform(TransformExecutionRequest(target=target, sources=[source], plan=proposal.plan.model_copy(update={"mappings": [mapping]})))
    except TransformValidationError as exc:
        return "incompatible" in str(exc)
    return False


def _suspicious_join_refused():
    request, _ = _ready_request()
    request = request.model_copy(update={"sources": [request.sources[0], _customers(duplicate=True)]})
    try:
        execute_transform(request)
    except TransformValidationError as exc:
        return "multiply rows" in str(exc)
    return False


def _unsupported_policy_refused(request):
    request = request.model_copy(update={"policy_evidence": [_evidence("This policy contains no supported commission percentage.")]})
    try:
        execute_transform(request)
    except TransformValidationError as exc:
        return "exactly one unambiguous" in str(exc)
    return False


def main() -> None:
    parser = argparse.ArgumentParser(description="Run controlled transform-to-template evaluations.")
    parser.add_argument("path", type=Path, help="Path to transform evaluation case definitions.")
    args = parser.parse_args()
    result = evaluate_template_transform_cases(args.path)
    print(json.dumps(result, indent=2))
    if result["failed_count"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
