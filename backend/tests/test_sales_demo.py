"""Known-ground-truth evaluation for the north-star August sales demo."""

from pathlib import Path
import base64
from io import BytesIO
from uuid import UUID

import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.models.grades import PolicyEvidence
from app.models.retrieval import SourceReference
from app.services.sales_report import SalesReportError, build_august_sales_report
from app.agent.models import AnswerClaim, Complete, ToolCall
from app.agent.orchestrator import AgentOrchestrator
from app.agent.providers import FakeModelProvider
from app.agent.tools import ToolRegistry, sales_report_tool
from app.agent.verification import EvidenceVerifier
from app.models.workflows import WorkflowCreate, WorkflowStep
from app.services.workflows import InMemoryWorkflowRepository, WorkflowService
from app.services.artifacts import InMemoryArtifactRepository
from app.main import app
from app.config import get_settings
from app.dependencies import _demo_repository, _deterministic_provider


ROOT = Path(__file__).parents[2]


def _evidence(text="Monthly commission policy: Salespeople earn a commission rate of 5% of completed net sales after discounts."):
    return [PolicyEvidence(text=text, source=SourceReference(document_id=UUID(int=1), filename="commission_policy.pdf", page_number=1, chunk_id=UUID(int=2)))]


def test_north_star_sales_report_known_outputs_and_provenance() -> None:
    transactions = pd.read_csv(ROOT / "sample_data" / "august_transactions.csv")
    original = transactions.copy(deep=True)
    result = build_august_sales_report(
        transactions,
        pd.read_csv(ROOT / "sample_data" / "sales_customers.csv"),
        pd.read_csv(ROOT / "sample_data" / "sales_targets.csv"),
        _evidence(),
    )

    pd.testing.assert_frame_equal(transactions, original)
    assert result.cleaned_transactions["transaction_id"].tolist() == ["T1001", "T1002", "T1003", "T1005"]
    regions = result.regional_performance.set_index("region")
    assert regions.loc["North", "net_sales"] == 1350
    assert regions.loc["North", "variance"] == -650
    assert regions.loc["South", "net_sales"] == 800
    assert regions.loc["South", "variance"] == -200
    assert regions.loc["West", "variance"] == -500
    commissions = result.commissions.set_index("salesperson")
    assert commissions.loc["Alice", "commission"] == 67.5
    assert commissions.loc["Bob", "commission"] == 50
    assert result.join_diagnostics.left_unmatched_rows == 1
    assert result.citations[0].filename == "commission_policy.pdf"
    assert result.artifact.content.startswith(b"PK")
    assert result.artifact.provenance["commission_rate"].startswith("5%")
    assert result.data_quality["original_transaction_rows"] == 6
    assert result.data_quality["duplicate_transaction_rows_removed"] == 1
    assert result.target_join_diagnostics["regions_missing_targets"] == 1
    assert result.verification_facts["total.net_sales"] == 2350
    assert result.verification_facts["total.commission"] == 117.5
    assert result.warnings
    assert result.reporting_period == "August 2026"
    workbook = load_workbook(BytesIO(result.artifact.content), data_only=False)
    assert workbook.sheetnames == [
        "Executive Summary",
        "Regional Performance",
        "Salesperson Performance",
        "Cleaned Transactions",
        "Data Quality",
        "Provenance & Sources",
    ]
    assert workbook["Cleaned Transactions"].max_row == 5
    assert workbook["Regional Performance"]["A2"].value == "North"
    assert workbook["Regional Performance"]["E2"].value == -650
    assert workbook["Salesperson Performance"]["D2"].value == 67.5
    assert len(workbook["Executive Summary"]._charts) == 1
    assert len(workbook["Regional Performance"]._charts) == 1
    assert len(workbook["Salesperson Performance"]._charts) == 1


def test_sales_report_supports_one_new_month_and_refuses_mixed_periods() -> None:
    september = pd.read_csv(ROOT / "sample_data" / "september_transactions.csv")
    customers = pd.read_csv(ROOT / "sample_data" / "sales_customers.csv")
    targets = pd.read_csv(ROOT / "sample_data" / "sales_targets.csv")

    result = build_august_sales_report(september, customers, targets, _evidence())

    assert result.reporting_period == "September 2026"
    assert result.verification_facts["total.net_sales"] == 3030
    assert result.artifact.filename == "september_sales_management_report.xlsx"
    workbook = load_workbook(BytesIO(result.artifact.content), data_only=False)
    assert workbook["Executive Summary"]["B2"].value == "Completed September 2026 net sales after discounts"

    mixed = pd.concat(
        [september, pd.read_csv(ROOT / "sample_data" / "august_transactions.csv").iloc[[0]]],
        ignore_index=True,
    )
    try:
        build_august_sales_report(mixed, customers, targets, _evidence())
    except SalesReportError as exc:
        assert "multiple reporting months" in str(exc)
    else:
        raise AssertionError("A report spanning multiple months must fail closed.")


def test_sales_report_refuses_missing_or_conflicting_policy() -> None:
    frames = (
        pd.read_csv(ROOT / "sample_data" / "august_transactions.csv"),
        pd.read_csv(ROOT / "sample_data" / "sales_customers.csv"),
        pd.read_csv(ROOT / "sample_data" / "sales_targets.csv"),
    )
    for evidence, message in [([], "does not establish"), (_evidence() + _evidence("Commission rate is 7%."), "conflicting")]:
        try:
            build_august_sales_report(*frames, evidence)
        except SalesReportError as exc:
            assert message in str(exc)
        else:
            raise AssertionError("Invalid policy evidence must fail.")


def test_sales_report_tool_is_verified_traced_and_rerunnable() -> None:
    def resource(name):
        return {"filename": name, "content_base64": base64.b64encode((ROOT / "sample_data" / name).read_bytes()).decode()}

    evidence = [item.model_dump(mode="json") for item in _evidence()]
    arguments = {
        "transactions": resource("august_transactions.csv"),
        "customers": resource("sales_customers.csv"),
        "targets": resource("sales_targets.csv"),
        "policy_evidence": evidence,
    }
    provider = FakeModelProvider([
        ToolCall(tool="sales.august_report", arguments=arguments),
        Complete(
            answer="North underperformed target by 650 and Alice earned 67.5 in commission.",
            claims=[
                AnswerClaim(text="North underperformed by 650", kind="numeric", value=650, evidence_keys=["regional.North.underperformance"]),
                AnswerClaim(text="Alice commission is 67.5", kind="numeric", value=67.5, evidence_keys=["commission.Alice"], source_ids=[str(UUID(int=2))]),
                AnswerClaim(text="Commission rate is policy-grounded", kind="document", source_ids=[str(UUID(int=2))]),
            ],
        ),
    ])
    artifact_repository = InMemoryArtifactRepository()
    registry = ToolRegistry([sales_report_tool(artifact_repository)])
    execution = AgentOrchestrator(provider, registry, EvidenceVerifier()).execute("Prepare the August sales report")
    assert execution.status == "completed"
    assert execution.verification.status == "verified_with_warnings"
    assert execution.trace[0].artifact_ids
    assert execution.trace[0].source_ids == [str(UUID(int=2))]

    workflows = WorkflowService(InMemoryWorkflowRepository(), registry)
    workflow = workflows.create(WorkflowCreate(name="August sales", source_task_id=execution.task_id, steps=[WorkflowStep(tool="sales.august_report", arguments=arguments)]))
    rerun = workflows.rerun(workflow.workflow_id, {})
    assert rerun.status == "completed"
    assert rerun.observations[0].result["regional_performance"]["rows"][0]["region"] == "North"

    previous_repository = app.state.artifact_repository
    app.state.artifact_repository = artifact_repository
    try:
        download = TestClient(app).get(f"/artifacts/{execution.trace[0].artifact_ids[0]}")
    finally:
        app.state.artifact_repository = previous_repository
    assert download.status_code == 200
    assert download.content.startswith(b"PK")
    assert download.headers["content-disposition"] == 'attachment; filename="august_sales_management_report.xlsx"'


def test_august_sales_demo_runs_through_public_http_and_downloads_workbook(monkeypatch) -> None:
    monkeypatch.setenv("APP_MODE", "demo")
    get_settings.cache_clear()
    _demo_repository.cache_clear()
    _deterministic_provider.cache_clear()
    client = TestClient(app)
    files = {
        "transactions": ("august_transactions.csv", (ROOT / "sample_data" / "august_transactions.csv").read_bytes(), "text/csv"),
        "customers": ("sales_customers.csv", (ROOT / "sample_data" / "sales_customers.csv").read_bytes(), "text/csv"),
        "targets": ("sales_targets.csv", (ROOT / "sample_data" / "sales_targets.csv").read_bytes(), "text/csv"),
        "policy": ("commission_policy.pdf", (ROOT / "sample_data" / "commission_policy.pdf").read_bytes(), "application/pdf"),
    }
    try:
        response = client.post("/sales/reports/august", files=files)
    finally:
        get_settings.cache_clear()
        _demo_repository.cache_clear()
        _deterministic_provider.cache_clear()

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "completed"
    assert "North has the largest target shortfall at 650.00" in body["answer"]
    assert [step["requested_tool"] for step in body["trace"]] == [
        "resource.list",
        "dataset.inspect",
        "dataset.inspect",
        "dataset.inspect",
        "document.search",
        "sales.north_star_report",
    ]
    assert body["verification"]["status"] == "verified_with_warnings"
    assert body["citations"][0]["filename"] == "commission_policy.pdf"
    assert body["artifacts"][0]["download_url"].startswith("/artifacts/")
    assert body["saved_workflow"]["name"] == "Monthly sales management report"
    assert [stage["name"] for stage in body["stages"]] == [
        "Goal",
        "Plan",
        "Inspect data",
        "Inspect data",
        "Inspect data",
        "Retrieve policy",
        "Clean",
        "Join",
        "Analyze",
        "Calculate commissions",
        "Generate charts and workbook",
        "Save workflow",
        "Verify",
        "Complete",
    ]
    assert body["warnings"]

    download = client.get(body["artifacts"][0]["download_url"])
    assert download.status_code == 200
    workbook = load_workbook(BytesIO(download.content), data_only=False)
    assert workbook["Executive Summary"]["B7"].value == 2350
    assert workbook["Executive Summary"]["B12"].value == 117.5

    workflow_id = body["saved_workflow"]["workflow_id"]
    saved = client.get(f"/workflows/{workflow_id}")
    rerun = client.post(f"/workflows/{workflow_id}/runs", json={"step_overrides": {}})
    compatible = client.post(
        f"/workflows/{workflow_id}/runs",
        json={
            "step_overrides": {
                "1": {
                    "transactions": {
                        "filename": "september_transactions.csv",
                        "content_base64": base64.b64encode(
                            (ROOT / "sample_data" / "september_transactions.csv").read_bytes()
                        ).decode(),
                    }
                }
            }
        },
    )
    drift = client.post(
        f"/workflows/{workflow_id}/runs",
        json={
            "step_overrides": {
                "1": {
                    "transactions": {
                        "filename": "incompatible_transactions.csv",
                        "content_base64": base64.b64encode(
                            (ROOT / "sample_data" / "incompatible_transactions.csv").read_bytes()
                        ).decode(),
                    }
                }
            }
        },
    )
    assert saved.status_code == 200
    assert saved.json()["steps"][0]["expected_schemas"]["transactions"] == [
        "transaction_id", "date", "salesperson", "customer_id", "amount", "discount", "status"
    ]
    assert rerun.status_code == 200
    assert rerun.json()["status"] == "completed"
    assert compatible.status_code == 200
    assert compatible.json()["status"] == "completed"
    compatible_total = next(
        item for item in compatible.json()["facts"] if item["key"] == "total.net_sales"
    )
    assert compatible_total["value"] == 3030
    comparison = client.post(
        "/workflow-runs/compare",
        json={
            "previous_run_id": rerun.json()["run_id"],
            "current_run_id": compatible.json()["run_id"],
        },
    )
    assert comparison.status_code == 200
    total_change = next(
        item for item in comparison.json()["metrics"] if item["label"] == "total.net_sales"
    )
    assert total_change["absolute_change"] == 680
    assert drift.status_code == 200
    assert drift.json()["status"] == "failed"
    assert "Schema drift detected for transactions" in drift.json()["error"]


def test_sales_report_rejects_conflicting_join_keys_and_preserves_missing_targets_as_warnings() -> None:
    transactions = pd.read_csv(ROOT / "sample_data" / "august_transactions.csv")
    customers = pd.read_csv(ROOT / "sample_data" / "sales_customers.csv")
    targets = pd.read_csv(ROOT / "sample_data" / "sales_targets.csv")
    conflicting_customers = pd.concat(
        [customers, pd.DataFrame([{"customer_id": "C001", "customer_name": "Conflict", "region": "South"}])],
        ignore_index=True,
    )

    try:
        build_august_sales_report(transactions, conflicting_customers, targets, _evidence())
    except SalesReportError as exc:
        assert "conflicting regions" in str(exc)
    else:
        raise AssertionError("Conflicting customer-to-region mappings must fail closed.")

    report = build_august_sales_report(transactions, customers, targets, _evidence())
    assert any("no target" in warning for warning in report.warnings)
    unassigned = report.regional_performance.set_index("region").loc["Unassigned"]
    assert unassigned["net_sales"] == 200
    assert pd.isna(unassigned["target"])


def test_general_agent_endpoint_runs_sales_from_schema_identified_resources(monkeypatch) -> None:
    monkeypatch.setenv("APP_MODE", "demo")
    get_settings.cache_clear()
    _demo_repository.cache_clear()
    _deterministic_provider.cache_clear()
    client = TestClient(app)
    policy = client.post(
        "/documents",
        files={
            "file": (
                "rules.pdf",
                (ROOT / "sample_data" / "commission_policy.pdf").read_bytes(),
                "application/pdf",
            )
        },
    )
    assert policy.status_code == 201

    def encoded(filename: str, bound_name: str) -> dict:
        return {
            "filename": bound_name,
            "content_base64": base64.b64encode((ROOT / "sample_data" / filename).read_bytes()).decode(),
        }

    payload = {
        "goal": "Prepare a verified August management report without changing the source files.",
        "max_iterations": 12,
        "resources": {
            "datasets": [
                encoded("august_transactions.csv", "input_a.csv"),
                encoded("sales_customers.csv", "input_b.csv"),
                encoded("sales_targets.csv", "input_c.csv"),
            ],
            "document_ids": [policy.json()["document_id"]],
        },
    }
    try:
        response = client.post("/agent/tasks", json=payload)
    finally:
        get_settings.cache_clear()
        _demo_repository.cache_clear()
        _deterministic_provider.cache_clear()

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "completed"
    assert body["trace"][-1]["requested_tool"] == "sales.north_star_report"
    assert body["verification"]["status"] == "verified_with_warnings"
    assert body["saved_workflow"] is not None
