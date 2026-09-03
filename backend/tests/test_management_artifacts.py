"""First-class artifact metadata and management workbook tests."""

from io import BytesIO
from uuid import uuid4

import pandas as pd
from openpyxl import load_workbook

from app.services.artifacts import InMemoryArtifactRepository, generate_management_workbook


def test_management_workbook_has_metadata_provenance_chart_and_safe_cells() -> None:
    task_id = uuid4()
    workflow_id = uuid4()
    frame = pd.DataFrame({"region": ["North", "=CMD()"], "revenue": [100, 50]})
    artifact = generate_management_workbook(
        frame, "August sales.csv", task_id=task_id, workflow_id=workflow_id,
        provenance={"source": "transactions.csv", "calculation": "deterministic"},
    )
    repository = InMemoryArtifactRepository()
    repository.save(artifact)
    restored = repository.get(artifact.artifact_id)

    assert restored is artifact
    assert artifact.filename == "August_sales_management_report.xlsx"
    assert artifact.producing_task_id == task_id
    assert artifact.producing_workflow_id == workflow_id
    workbook = load_workbook(BytesIO(artifact.content), data_only=False)
    assert workbook.sheetnames == ["Data", "Summary", "Provenance"]
    assert workbook["Data"]["A3"].value == "'=CMD()"
    assert len(workbook["Summary"]._charts) == 1
