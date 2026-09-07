import base64
from io import BytesIO
import json
from pathlib import Path
from uuid import UUID, uuid4

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.agent.models import AgentDatasetResource, AgentTaskResources
from app.agent.tools import ToolRegistry, general_task_tools, template_transform_tool
from app.models.grades import PolicyEvidence
from app.models.retrieval import SourceReference
from app.models.template_transforms import (
    DerivationRule,
    FieldMapping,
    FieldReference,
    FilePayload,
    JoinRule,
    SourcePayload,
    TransformExecutionRequest,
    TransformProposalRequest,
)
from app.models.workflows import WorkflowCreate, WorkflowStep
from app.services.artifacts import InMemoryArtifactRepository
from app.services.template_transforms import (
    ClarificationRequiredError,
    TemplateInspectionError,
    TransformValidationError,
    execute_transform,
    inspect_template,
    propose_transform,
)
from app.services.workflows import InMemoryWorkflowRepository, WorkflowService
from app.main import app
from app.config import get_settings


ROOT = Path(__file__).parents[2]


def _csv(rows: list[dict]) -> bytes:
    return pd.DataFrame(rows).to_csv(index=False).encode()


def _xlsx_template(*, include_approval: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Submission"
    headers = ["Order ID", "Customer Name", "Region", "Net Sales", "Commission", "Total"]
    if include_approval:
        headers.insert(-1, "Manager Approval Code")
    sheet.append(headers)
    values = [1, "Example", "North", 100.0, 5.0, "=D2+E2"]
    if include_approval:
        values.insert(-1, "APPROVED")
        values[-1] = "=D2+E2"
    sheet.append(values)
    sheet["D2"].number_format = "$#,##0.00"
    sheet["E2"].number_format = "$#,##0.00"
    notes = workbook.create_sheet("Instructions")
    notes["A1"] = "Do not remove this sheet."
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _source(filename: str, role: str, content: bytes) -> SourcePayload:
    return SourcePayload(filename=filename, role=role, content_base64=base64.b64encode(content).decode())


def _target(content: bytes) -> FilePayload:
    return FilePayload(filename="required_template.xlsx", content_base64=base64.b64encode(content).decode(), sheet="Submission")


def _orders() -> SourcePayload:
    return _source(
        "raw_orders.csv",
        "orders",
        _csv(
            [
                {"order_id": 101, "customer_id": "C1", "gross_sales": "$1,000.00", "returns": "100", "eligible_sales": 900},
                {"order_id": 102, "customer_id": "C2", "gross_sales": "$500.00", "returns": "50", "eligible_sales": 450},
            ]
        ),
    )


def _customers(rows: list[dict] | None = None) -> SourcePayload:
    return _source(
        "customer_master.csv",
        "customers",
        _csv(rows or [{"cust_id": "C1", "customer_name": "Alice", "territory": "North"}, {"cust_id": "C2", "customer_name": "=HYPERLINK(\"x\")", "territory": "South"}]),
    )


def _evidence(text: str = "The commission policy rate is 5% of eligible sales.") -> PolicyEvidence:
    return PolicyEvidence(
        text=text,
        source=SourceReference(document_id=uuid4(), filename="reporting_policy.pdf", page_number=1, chunk_id=uuid4()),
    )


def _ready_request(customers: SourcePayload | None = None, evidence: list[PolicyEvidence] | None = None):
    template = _xlsx_template()
    sources = [_orders(), customers or _customers()]
    proposal = propose_transform(TransformProposalRequest(target=_target(template), sources=sources))
    mappings = []
    for mapping in proposal.plan.mappings:
        if mapping.target_field in {"Net Sales", "Commission"}:
            mappings.append(mapping)
        else:
            mappings.append(mapping.model_copy(update={"status": "confirmed"}))
    plan = proposal.plan.model_copy(
        update={
            "mappings": mappings,
            "joins": [JoinRule(right_role="customers", left_on=["orders.customer_id"], right_on=["cust_id"])],
            "derivations": [
                DerivationRule(
                    target_field="Net Sales",
                    operation="subtract",
                    inputs=[
                        FieldReference(source_role="orders", source_field="gross_sales", transformation="normalize_currency"),
                        FieldReference(source_role="orders", source_field="returns"),
                    ],
                ),
                DerivationRule(
                    target_field="Commission",
                    operation="policy_multiply",
                    inputs=[FieldReference(source_role="target", source_field="Net Sales")],
                    policy_query="commission policy rate",
                ),
            ],
            "unique_fields": ["Order ID"],
        }
    )
    return TransformExecutionRequest(target=_target(template), sources=sources, plan=plan, policy_evidence=evidence or [_evidence()])


def test_template_inspection_preserves_structure_and_fingerprints_formulas():
    result = inspect_template("required_template.xlsx", _xlsx_template(), "Submission")

    assert result.headers == ["Order ID", "Customer Name", "Region", "Net Sales", "Commission", "Total"]
    assert [sheet.name for sheet in result.sheets] == ["Submission", "Instructions"]
    assert result.columns[-1].formula == "=D2+E2"
    assert result.columns[-1].formula_row == 2
    assert result.columns[3].number_format == "$#,##0.00"
    assert len(result.fingerprint) == 64
    assert result.has_macros is False


def test_proposal_uses_deterministic_mapping_order_and_requires_missing_business_data():
    template = _xlsx_template(include_approval=True)
    proposal = propose_transform(TransformProposalRequest(target=_target(template), sources=[_orders(), _customers()]))
    by_target = {item.target_field: item for item in proposal.plan.mappings}

    assert by_target["Order ID"].mapping_type == "normalized"
    assert by_target["Customer Name"].source_role == "customers"
    assert by_target["Region"].mapping_type == "documented_alias"
    assert by_target["Total"].mapping_type == "template_formula"
    assert any(item.code == "missing_required_field" and item.target_field == "Manager Approval Code" for item in proposal.clarifications)
    assert proposal.status == "clarification_required"


def test_proposal_reports_ambiguous_aliases_instead_of_guessing():
    template = b"Region\n"
    first = _source("territories.csv", "territories", _csv([{"territory": "North"}]))
    second = _source("areas.csv", "areas", _csv([{"sales_area": "N"}]))

    proposal = propose_transform(
        TransformProposalRequest(
            target=FilePayload(filename="target.csv", content_base64=base64.b64encode(template).decode()),
            sources=[first, second],
        )
    )

    assert proposal.status == "clarification_required"
    assert proposal.clarifications[0].code == "ambiguous_mapping"
    assert proposal.clarifications[0].target_field == "Region"


def test_invalid_explicit_mapping_returns_structured_incompatible_clarification():
    target = FilePayload(filename="target.csv", content_base64=base64.b64encode(b"Region\n").decode())
    source = _source("territories.csv", "territories", _csv([{"territory": "North"}]))

    proposal = propose_transform(
        TransformProposalRequest(target=target, sources=[source], explicit_mappings={"Region": "unknown.field"})
    )

    assert proposal.status == "clarification_required"
    assert proposal.plan.mappings[0].status == "incompatible"
    assert proposal.clarifications[0].code == "incompatible_type"


def test_execute_writes_exact_template_preserves_other_sheets_formulas_and_neutralizes_injection():
    executed = execute_transform(_ready_request())
    workbook = load_workbook(BytesIO(executed.artifact.content), data_only=False)
    sheet = workbook["Submission"]

    assert executed.validation.status == "completed"
    assert workbook.sheetnames == ["Submission", "Instructions"]
    assert [sheet.cell(1, column).value for column in range(1, 7)] == ["Order ID", "Customer Name", "Region", "Net Sales", "Commission", "Total"]
    assert sheet["D2"].value == 900
    assert sheet["E2"].value == 45
    assert sheet["F2"].value == "=D2+E2"
    assert sheet["F3"].value == "=D3+E3"
    assert sheet["B3"].value.startswith("'=")
    assert workbook["Instructions"]["A1"].value == "Do not remove this sheet."
    assert any(item.target_field == "Commission" and item.policy_evidence_ids for item in executed.provenance)
    assert all(check.passed for check in executed.validation.checks)


def test_required_template_formula_fields_are_satisfied_by_trusted_formulas():
    request = _ready_request()
    plan = request.plan.model_copy(update={"required_fields": [*request.plan.required_fields, "Total"]})

    executed = execute_transform(request.model_copy(update={"plan": plan}))
    workbook = load_workbook(BytesIO(executed.artifact.content), data_only=False)

    assert executed.validation.status == "completed"
    assert workbook["Submission"]["F2"].value == "=D2+E2"
    assert workbook["Submission"]["F3"].value == "=D3+E3"
    assert all(check.passed for check in executed.validation.checks if check.name.startswith("required:"))


def test_execute_fails_closed_on_unresolved_required_field():
    template = _xlsx_template(include_approval=True)
    proposal = propose_transform(TransformProposalRequest(target=_target(template), sources=[_orders(), _customers()]))

    with pytest.raises(ClarificationRequiredError) as exc:
        execute_transform(TransformExecutionRequest(target=_target(template), sources=[_orders(), _customers()], plan=proposal.plan))

    assert any(item.target_field == "Manager Approval Code" for item in exc.value.clarifications)


def test_execute_treats_whitespace_only_required_values_as_missing():
    target = FilePayload(filename="target.csv", content_base64=base64.b64encode(b"Name,ID\n").decode())
    source = _source("source.csv", "source", b"Name,ID\n   ,1\n")
    proposal = propose_transform(TransformProposalRequest(target=target, sources=[source]))

    with pytest.raises(TransformValidationError, match="Required field 'Name'"):
        execute_transform(TransformExecutionRequest(target=target, sources=[source], plan=proposal.plan))


def test_execute_blocks_row_multiplying_join():
    customers = _customers(
        [
            {"cust_id": "C1", "customer_name": "Alice", "territory": "North"},
            {"cust_id": "C1", "customer_name": "Alicia", "territory": "West"},
            {"cust_id": "C2", "customer_name": "Bob", "territory": "South"},
        ]
    )
    request = _ready_request(customers=customers)

    with pytest.raises(TransformValidationError, match="multiply rows"):
        execute_transform(request)


def test_policy_derivation_refuses_missing_or_conflicting_rates():
    missing = _ready_request(evidence=[])
    missing = missing.model_copy(update={"policy_evidence": []})
    with pytest.raises(TransformValidationError, match="requires retrieved policy evidence"):
        execute_transform(missing)

    conflicting = _ready_request(evidence=[_evidence("Commission is 5%. Another applicable rate is 7%.")])
    with pytest.raises(TransformValidationError, match="exactly one unambiguous"):
        execute_transform(conflicting)


def test_workflow_rerun_reuses_confirmed_plan_and_detects_schema_and_template_drift():
    request = _ready_request()
    artifacts = InMemoryArtifactRepository()
    service = WorkflowService(InMemoryWorkflowRepository(), ToolRegistry([template_transform_tool(artifacts)]), artifacts)
    workflow = service.create(
        WorkflowCreate(name="Monthly submission", steps=[WorkflowStep(tool="template.transform", arguments=request.model_dump(mode="json"))])
    )

    completed = service.rerun(workflow.workflow_id, {})
    assert completed.status == "completed"
    assert completed.observations[0].artifact_ids

    drifted_source = _source("raw_orders.csv", "orders", _csv([{"order_id": 1, "customer_id": "C1", "eligible_sales": 10, "returns": 1, "new_column": "drift"}]))
    failed = service.rerun(workflow.workflow_id, {1: {"sources": [drifted_source.model_dump(mode="json"), request.sources[1].model_dump(mode="json")]}})
    assert failed.status == "failed"
    assert "Schema drift" in failed.error

    other_template = FilePayload(filename="required_template.xlsx", content_base64=base64.b64encode(_xlsx_template(include_approval=True)).decode(), sheet="Submission")
    template_failed = service.rerun(workflow.workflow_id, {1: {"target": other_template.model_dump(mode="json")}})
    assert template_failed.status == "failed"
    assert "Template drift" in template_failed.error

    plan_override = service.rerun(workflow.workflow_id, {1: {"plan": request.plan.model_dump(mode="json")}})
    assert plan_override.status == "failed"
    assert "pinned" in plan_override.error


def test_macro_enabled_and_unsafe_workbook_extension_is_rejected():
    with pytest.raises(TemplateInspectionError, match="macro-enabled"):
        inspect_template("target.xlsm", _xlsx_template())

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["External"])
    sheet.append(['=WEBSERVICE("https://example.invalid")'])
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(TemplateInspectionError, match="unsupported external"):
        inspect_template("target.xlsx", output.getvalue())

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Local table formula"])
    sheet.append(["=[@[Net Sales]]+[@Commission]"])
    output = BytesIO()
    workbook.save(output)
    inspected = inspect_template("target.xlsx", output.getvalue())
    assert inspected.columns[0].formula == "=[@[Net Sales]]+[@Commission]"


def test_public_api_proposes_executes_persists_and_downloads_exact_csv():
    artifacts = InMemoryArtifactRepository()
    workflows = WorkflowService(InMemoryWorkflowRepository(), ToolRegistry([template_transform_tool(artifacts)]), artifacts)
    previous_artifacts = app.state.artifact_repository
    previous_workflows = app.state.workflow_service
    app.state.artifact_repository = artifacts
    app.state.workflow_service = workflows
    client = TestClient(app)
    try:
        files = [
            ("target", ("target.csv", b"Order ID\n", "text/csv")),
            ("sources", ("orders.csv", b"order_id\n101\n102\n", "text/csv")),
        ]
        proposal_response = client.post(
            "/template-transforms/proposals",
            files=files,
            data={"source_roles": '["orders"]'},
        )
        assert proposal_response.status_code == 200
        proposal = proposal_response.json()
        assert proposal["status"] == "ready"

        execution_response = client.post(
            "/template-transforms/executions",
            files=files,
            data={"source_roles": '["orders"]', "plan": json.dumps(proposal["plan"])},
        )
        assert execution_response.status_code == 200
        body = execution_response.json()
        assert body["status"] == "completed"
        assert body["saved_workflow"]["workflow_id"]
        artifact_response = client.get(body["artifact"]["download_url"])
        assert artifact_response.status_code == 200
        assert artifact_response.content.replace(b"\r\n", b"\n") == b"Order ID\n101\n102\n"
    finally:
        app.state.artifact_repository = previous_artifacts
        app.state.workflow_service = previous_workflows


def test_canonical_demo_runs_end_to_end_through_api_with_retrieved_policy(monkeypatch):
    monkeypatch.setenv("APP_MODE", "demo")
    get_settings.cache_clear()
    settings = get_settings()
    app.dependency_overrides[get_settings] = lambda: settings
    artifacts = InMemoryArtifactRepository()
    workflows = WorkflowService(InMemoryWorkflowRepository(), ToolRegistry([template_transform_tool(artifacts)]), artifacts)
    previous_artifacts = app.state.artifact_repository
    previous_workflows = app.state.workflow_service
    app.state.artifact_repository = artifacts
    app.state.workflow_service = workflows
    client = TestClient(app)
    sample = ROOT / "sample_data"
    proposal_files = [
        ("target", ("required_template.xlsx", (sample / "required_template.xlsx").read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
        ("sources", ("raw_orders.xlsx", (sample / "raw_orders.xlsx").read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
        ("sources", ("customer_master.csv", (sample / "customer_master.csv").read_bytes(), "text/csv")),
    ]
    try:
        proposal_response = client.post(
            "/template-transforms/proposals",
            files=proposal_files,
            data={"source_roles": '["orders","customers"]'},
        )
        assert proposal_response.status_code == 200
        plan = proposal_response.json()["plan"]
        plan["joins"] = [{
            "right_role": "customers", "left_on": ["orders.customer_id"], "right_on": ["cust_id"], "how": "left",
            "block_many_to_many": True, "block_row_multiplication": True, "max_unmatched_left_percentage": 0,
        }]
        plan["derivations"] = [
            {
                "target_field": "Net Sales", "operation": "subtract", "separator": " ", "constant": None, "policy_query": None,
                "inputs": [
                    {"source_role": "orders", "source_field": "gross_sales", "transformation": "normalize_currency"},
                    {"source_role": "orders", "source_field": "returns", "transformation": "none"},
                ],
            },
            {
                "target_field": "Commission", "operation": "policy_multiply", "separator": " ", "constant": None,
                "policy_query": "commission policy rate",
                "inputs": [{"source_role": "target", "source_field": "Net Sales", "transformation": "none"}],
            },
        ]
        plan["unique_fields"] = ["Order ID"]
        execution_files = [
            *proposal_files,
            ("policy", ("reporting_policy.pdf", (sample / "reporting_policy.pdf").read_bytes(), "application/pdf")),
        ]
        response = client.post(
            "/template-transforms/executions",
            files=execution_files,
            data={"source_roles": '["orders","customers"]', "plan": json.dumps(plan)},
        )
        assert response.status_code == 200
        body = response.json()
        assert body["status"] == "completed"
        assert body["saved_workflow"]["workflow_id"]
        artifact = artifacts.get(UUID(body["artifact"]["artifact_id"]))
        workbook = load_workbook(BytesIO(artifact.content), data_only=False)
        assert workbook.sheetnames == ["Monthly Submission", "Instructions"]
        assert workbook["Monthly Submission"]["D2"].value == 900
        assert workbook["Monthly Submission"]["E2"].value == 45
        assert workbook["Monthly Submission"]["B4"].value.startswith("'=")
        assert any(item["policy_evidence_ids"] for item in body["provenance"] if item["target_field"] == "Commission")
    finally:
        app.state.artifact_repository = previous_artifacts
        app.state.workflow_service = previous_workflows
        app.dependency_overrides.clear()
        get_settings.cache_clear()


def test_task_scoped_agent_tools_can_propose_execute_and_save_template_workflow():
    artifacts = InMemoryArtifactRepository()
    workflows = WorkflowService(InMemoryWorkflowRepository(), ToolRegistry([template_transform_tool(artifacts)]), artifacts)
    resources = AgentTaskResources(
        datasets=[
            AgentDatasetResource(filename="target.csv", content_base64=base64.b64encode(b"Order ID\n").decode()),
            AgentDatasetResource(filename="orders.csv", content_base64=base64.b64encode(b"order_id\n101\n").decode()),
        ]
    )
    registry = ToolRegistry(general_task_tools(None, resources, artifacts, workflows))
    for name in ("target.csv", "orders.csv"):
        tool = registry.get("dataset.inspect")
        tool.handler(tool.input_model.model_validate({"dataset": name}))

    propose_tool = registry.get("template.propose")
    proposal_observation = propose_tool.handler(
        propose_tool.input_model.model_validate(
            {"target_dataset": "target.csv", "sources": [{"role": "orders", "dataset": "orders.csv"}]}
        )
    )
    assert proposal_observation.result["status"] == "ready"

    execute_tool = registry.get("template.execute")
    execution = execute_tool.handler(
        execute_tool.input_model.model_validate(
            {
                "target_dataset": "target.csv",
                "sources": [{"role": "orders", "dataset": "orders.csv"}],
                "plan": proposal_observation.result["plan"],
            }
        )
    )
    assert execution.success is True
    assert execution.result["saved_workflow"]["workflow_id"]
    assert execution.artifact_ids
    saved = workflows.get(UUID(execution.result["saved_workflow"]["workflow_id"]))
    assert all(
        mapping["status"] == "confirmed"
        for mapping in saved.steps[0].arguments["plan"]["mappings"]
    )
