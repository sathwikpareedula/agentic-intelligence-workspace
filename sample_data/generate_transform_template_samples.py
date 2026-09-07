"""Generate deterministic binary fixtures for the transform-to-template demo."""

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


ROOT = Path(__file__).parent


def write_pdf(path: Path, text: str) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream.set_data(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    path.write_bytes(output.getvalue())


def write_orders(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Raw Orders"
    sheet.append(["order_id", "customer_id", "gross_sales", "returns", "eligible_sales"])
    sheet.append(["O-1001", "C001", "$1,000.00", 100, 900])
    sheet.append(["O-1002", "C002", "$750.00", 50, 700])
    sheet.append(["O-1003", "C003", "$500.00", 0, 500])
    workbook.save(path)


def write_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Monthly Submission"
    headers = ["Order ID", "Customer Name", "Region", "Net Sales", "Commission", "Total"]
    sheet.append(headers)
    sheet.append(["example", "Example Customer", "North", 100.0, 5.0, "=D2+E2"])
    fill = PatternFill("solid", fgColor="174C3C")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
    for cell in (sheet["D2"], sheet["E2"], sheet["F2"]):
        cell.number_format = "$#,##0.00"
    instructions = workbook.create_sheet("Instructions")
    instructions["A1"] = "Populate Monthly Submission without removing this sheet or changing column order."
    workbook.save(path)


if __name__ == "__main__":
    write_orders(ROOT / "raw_orders.xlsx")
    write_template(ROOT / "required_template.xlsx")
    write_pdf(
        ROOT / "reporting_policy.pdf",
        "Monthly reporting policy: Eligible net sales receive a commission rate of 5%. Use only completed supplied records.",
    )

